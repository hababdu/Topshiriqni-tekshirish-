"""Statistikani Excel (.xlsx) fayliga chiqarish."""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import Config
from .db import DB
from .timeutil import today, week_range

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="1F2937")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PERIODS = {
    "w0": "Shu hafta",
    "w1": "O'tgan hafta",
    "m": "Shu oy",
    "all": "Hammasi",
}


def period_range(key: str, tz: str) -> tuple[str, str, str]:
    """(boshlanish, tugash-eksklyuziv, fayl nomi uchun belgi)"""
    d = today(tz)
    if key == "w0":
        s, e = week_range(d, 0)
        return s, e, f"hafta_{s}"
    if key == "w1":
        s, e = week_range(d, 1)
        return s, e, f"hafta_{s}"
    if key == "m":
        first = d.replace(day=1)
        nxt = (first + timedelta(days=32)).replace(day=1)
        return first.isoformat(), nxt.isoformat(), f"oy_{first:%Y-%m}"
    return "2000-01-01", "2100-01-01", f"hammasi_{d.isoformat()}"


def _fmt_date(iso: str) -> str:
    try:
        return date.fromisoformat(iso).strftime("%d.%m.%Y")
    except ValueError:
        return iso


def _header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, min_w: int = 8, max_w: int = 45) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                longest = max(len(x) for x in str(c.value).split("\n"))
                widths[c.column] = max(widths.get(c.column, 0), longest)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def _score(v: float | None):
    return None if v is None else round(v, 1)


async def build_stats_xlsx(db: DB, cfg: Config, period: str) -> tuple[bytes, str]:
    start, end, tag = period_range(period, cfg.timezone)
    label = PERIODS.get(period, period)

    stats = await db.week_stats(start, end)  # hozirgi o'quvchilar bo'yicha reyting
    roster = await db.all_students()
    tasks = await db.tasks_in_range(start, end)
    smap = await db.score_map([t["id"] for t in tasks])
    log_rows = await db.submissions_in_range(start, end)
    by_user = {s["user_id"]: s for s in roster}

    wb = Workbook()

    # ---- 1-varaq: Reyting ----
    ws = wb.active
    ws.title = "Reyting"
    ws["A1"] = f"Statistika — {label} ({_fmt_date(start)} — {_fmt_date((date.fromisoformat(end) - timedelta(days=1)).isoformat()) if period != 'all' else 'bugun'})"
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, 3, ["№", "O'quvchi", "Topshirilgan vazifalar", "Yuborilgan ishlar", "Baholangan", "Jami ball", "O'rtacha ball"])
    n_tasks = len(tasks)
    for i, s in enumerate(stats, 1):
        done = sum(1 for t in tasks if t["id"] in smap.get(s["user_id"], {}))
        row = 3 + i
        vals = [
            i,
            s["full_name"],
            f"{done}/{n_tasks}" if n_tasks else "—",
            s["total"],
            s["graded"],
            s["points"],
            _score(s["avg_score"]),
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = BORDER
            if col != 2:
                c.alignment = Alignment(horizontal="center")
        if n_tasks:
            ws.cell(row=row, column=3).fill = GREEN if done == n_tasks else RED
    ws.freeze_panes = "A4"
    _autosize(ws)

    # ---- 2-varaq: Vazifalar holati (yashil/qizil) ----
    ws2 = wb.create_sheet("Vazifalar holati")
    heads = ["№", "O'quvchi"] + [
        f"{_fmt_date(t['week_start'])}\n{t['title']}" for t in tasks
    ] + ["Topshirgan", "Topshirmagan"]
    _header(ws2, 1, heads)
    ws2.row_dimensions[1].height = 45
    for i, s in enumerate(roster, 1):
        row = 1 + i
        done = smap.get(s["user_id"], {})
        ws2.cell(row=row, column=1, value=i).border = BORDER
        name = ws2.cell(row=row, column=2, value=s["full_name"] + ("" if s["started"] else "  ⚠"))
        name.border = BORDER
        n_ok = 0
        for j, t in enumerate(tasks):
            c = ws2.cell(row=row, column=3 + j)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
            if t["id"] in done:
                n_ok += 1
                c.value = _score(done[t["id"]]) if done[t["id"]] is not None else "✔"
                c.fill = GREEN
            else:
                c.value = "✘"
                c.fill = RED
        for k, v in enumerate((n_ok, len(tasks) - n_ok)):
            c = ws2.cell(row=row, column=3 + len(tasks) + k, value=v)
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
        ws2.cell(row=row, column=3 + len(tasks)).fill = GREEN
        ws2.cell(row=row, column=4 + len(tasks)).fill = RED if len(tasks) - n_ok else GREEN
    legend_row = len(roster) + 3
    ws2.cell(row=legend_row, column=2, value="Katak: ball — baholangan, ✔ — topshirilgan (hali baholanmagan), ✘ — topshirilmagan. ⚠ — botga /start bosmagan.")
    if not tasks:
        ws2.cell(row=2, column=3, value="Bu davrda vazifa yo'q")
    ws2.freeze_panes = "C2"
    _autosize(ws2, max_w=28)
    ws2.column_dimensions["B"].width = 32

    # ---- 3-varaq: Barcha ishlar ----
    ws3 = wb.create_sheet("Barcha ishlar")
    _header(ws3, 1, ["Sana", "O'quvchi", "Vazifa", "Turi", "Ball"])
    for i, r in enumerate(log_rows, 2):
        vals = [_fmt_date(r["submit_date"]), r["full_name"], r["task_title"], r["kind"], r["score"] if r["score"] is not None else "—"]
        for col, v in enumerate(vals, 1):
            ws3.cell(row=i, column=col, value=v).border = BORDER
    ws3.freeze_panes = "A2"
    _autosize(ws3)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"statistika_{tag}.xlsx"
